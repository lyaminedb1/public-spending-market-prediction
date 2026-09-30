"""Mise en forme ECE du Word produit par pandoc (Times New Roman 12, double interligne, marges 2,54 cm, titres, tableaux)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

import pathlib
B = pathlib.Path(__file__).resolve().parents[1] / "build"
d = Document(B / "brut.docx")
TNR = "Times New Roman"

def font(style, size=12, bold=None, italic=None):
    f = style.font; f.name = TNR; f.size = Pt(size); f.color.rgb = RGBColor(0, 0, 0)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), TNR)
    for t in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(t)) is not None: del rf.attrib[qn(t)]

# Marges 2,54 cm, A4
for s in d.sections:
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(s, m, Cm(2.54))
    # numéro de page en bas au centre
    p = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)
    r.font.name = TNR; r.font.size = Pt(11)

styles = d.styles
STY = {x.name: x for x in d.styles}
for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
    if name in [s.name for s in styles]:
        st = STY[name]; font(st, 12)
        pf = st.paragraph_format; pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.space_before = Pt(0); pf.space_after = Pt(6)
        if name != "Compact": pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for lvl, size, align, indent in ((1, 14, WD_ALIGN_PARAGRAPH.CENTER, 0), (2, 12, WD_ALIGN_PARAGRAPH.LEFT, 0), (3, 12, WD_ALIGN_PARAGRAPH.LEFT, 1.27)):
    st = STY[f"Heading {lvl}"]; font(st, size, bold=True, italic=False)
    pf = st.paragraph_format; pf.alignment = align; pf.left_indent = Cm(indent)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE; pf.space_before = Pt(18 if lvl < 3 else 12); pf.space_after = Pt(12 if lvl < 3 else 6)
    pf.keep_with_next = True
if "Bibliography" in [s.name for s in styles]:
    st = STY["Bibliography"]; font(st, 12)
    pf = st.paragraph_format; pf.line_spacing_rule = WD_LINE_SPACING.SINGLE; pf.space_after = Pt(8)
    pf.left_indent = Cm(1.27); pf.first_line_indent = Cm(-1.27); pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
for name in ("Hyperlink",):
    if name in [s.name for s in styles]: STY[name].font.name = TNR

# Titres de niveau 3 : terminés par un point (guide ECE)
for p in d.paragraphs:
    if p.style.name == "Heading 3" and p.runs:
        t = p.text.rstrip()
        if t and t[-1] not in ".?!:":
            p.runs[-1].text = p.runs[-1].text.rstrip() + "."
    # légendes et sources de tableaux : simple interligne
    if re.match(r"^(Tableau \d\.\d|Figure \d\.\d|Source :)", p.text):
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.keep_with_next = p.text.startswith("Tableau")
        if p.text.startswith("Figure"): p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6 if p.text.startswith("Tableau") else 2)
        p.paragraph_format.space_after = Pt(4 if p.text.startswith("Tableau") else 12)
        for r in p.runs: r.font.size = Pt(10 if p.text.startswith("Source") else 11)

# Tableaux : bordures, simple interligne, police 9-10
def borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "808080"); b.append(e)
    tblPr.append(b)
for tbl in d.tables:
    borders(tbl)
    ncol = len(tbl.columns)
    size = 9 if ncol >= 5 else 10
    for ri, row in enumerate(tbl.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                p.paragraph_format.space_after = Pt(2); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(size); r.font.name = TNR
                    if ri == 0: r.font.bold = True
            if ri == 0:
                tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
                sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), "E7E6E6"); tcPr.append(sh)
d.core_properties.author = "Abdellah Elyamine DALI BRAHAM"
d.core_properties.title = "Mémoire – Chapitres 1 à 4"
d.save(B / "Memoire_DaliBraham.docx")
print("ok")

# paragraphes contenant une équation : simple interligne (affichage LibreOffice)
d = Document(B / "Memoire_DaliBraham.docx")
for p in d.paragraphs:
    if p._p.xpath(".//m:oMathPara") or p._p.xpath(".//m:oMath"):
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)
d.save(B / "Memoire_DaliBraham.docx")
print("équations ok")
